"""The Voucher Configuration, as the voucher page and the import templates show it.

The server already enforced it on save and on import. What people saw did not:

  * the Receipt template offered all 275 ledgers on both sides, and its sample
    was one Credit line to "Customer A" dated 2023-12-31 - unbalanced, in a
    format and a year nothing else used;
  * the voucher page, redrawn after a rejected save, emptied every ledger
    dropdown - it was handed ledger records where it compares names;
  * the page knew the Receipt and Payment cash rules but not Contra's.
"""
import contextlib
import io
import os
import sys
import unittest

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), '..')))

with contextlib.redirect_stdout(io.StringIO()):
    import app as appmod
from accounting_app import models as M
from database.config import get_connection


def _admin_and_company():
    conn = get_connection()
    try:
        cur = conn.cursor()
        cur.execute("SELECT id FROM users WHERE username = 'admin'")
        user = cur.fetchone()
        cur.execute("SELECT company_id FROM ledgers GROUP BY company_id "
                    "ORDER BY COUNT(*) DESC LIMIT 1")
        company = cur.fetchone()
        return (user[0] if user else None), (company[0] if company else None)
    finally:
        conn.close()


class AllowedLedgerNamesTests(unittest.TestCase):
    """One answer to "which ledgers may this side use", matching the validator."""

    def setUp(self):
        _user, self.company = _admin_and_company()
        if not self.company:
            self.skipTest("needs a company with ledgers")

    def test_the_built_in_cash_rules_are_reported(self):
        for vt, side in (("Receipt", "Debit"), ("Payment", "Credit"),
                         ("Contra", "Debit"), ("Contra", "Credit")):
            names = M.allowed_ledger_names(vt, side, company_id=self.company)
            self.assertIsNotNone(names, (vt, side))

    def test_every_listed_ledger_passes_the_validator(self):
        # The template offers exactly what save and import accept.
        for vt, side in (("Receipt", "Debit"), ("Payment", "Credit"),
                         ("Contra", "Debit"), ("Contra", "Credit")):
            for name in M.allowed_ledger_names(vt, side, company_id=self.company) or []:
                ok, err = M.validate_voucher_ledger_groups(
                    vt, [{"ledger_name": name, "amount": 1, "type": side}],
                    company_id=self.company)
                self.assertTrue(ok, (vt, side, name, err))

    def test_a_ledger_left_off_the_list_is_refused_by_the_validator(self):
        allowed = set(M.allowed_ledger_names("Receipt", "Debit", company_id=self.company) or [])
        conn = get_connection()
        try:
            cur = conn.cursor()
            cur.execute("SELECT ledger_name FROM ledgers WHERE company_id = %s "
                        "AND COALESCE(is_active, 1) = 1", (self.company,))
            outside = [r[0] for r in cur.fetchall() if r[0] not in allowed]
        finally:
            conn.close()
        if not outside:
            self.skipTest("every ledger is a cash or bank ledger")
        ok, _err = M.validate_voucher_ledger_groups(
            "Receipt", [{"ledger_name": outside[0], "amount": 1, "type": "Debit"}],
            company_id=self.company)
        self.assertFalse(ok)

    def test_an_unrestricted_side_says_so(self):
        # Journal has no built-in rule; with no configuration, any ledger.
        from database.voucher_config_db import get_voucher_config
        if get_voucher_config("Journal", "Debit", company_id=self.company):
            self.skipTest("this company configures Journal")
        self.assertIsNone(M.allowed_ledger_names("Journal", "Debit", company_id=self.company))


class VoucherPageTests(unittest.TestCase):

    def test_a_redrawn_form_gets_names_not_records(self):
        user, company = _admin_and_company()
        if not company:
            self.skipTest("needs a company")
        from accounting_app.voucher_routes import _voucher_form_context
        with appmod.app.test_request_context("/"):
            from flask_login import login_user
            from accounting_app.models import User
            conn = get_connection()
            try:
                cur = conn.cursor()
                cur.execute("SELECT id, username, email, password_hash, is_admin "
                            "FROM users WHERE id = %s", (user,))
                row = cur.fetchone()
            finally:
                conn.close()
            login_user(User(*row))
            with contextlib.redirect_stdout(io.StringIO()):
                context = _voucher_form_context("Receipt", company)
        for key in ("allowed_ledgers_dr", "allowed_ledgers_cr"):
            self.assertTrue(context[key], key)
            self.assertTrue(all(isinstance(v, str) for v in context[key]), key)

    def test_the_page_applies_the_contra_rule(self):
        with io.open(os.path.join(os.path.dirname(__file__), "..", "templates",
                                  "voucher.html"), encoding="utf-8") as f:
            page = f.read()
        self.assertIn("|| voucherType === 'Contra') {", page)


class TemplateTests(unittest.TestCase):
    """The Receipt template, downloaded and read back."""

    @classmethod
    def setUpClass(cls):
        user, company = _admin_and_company()
        if not (user and company):
            raise unittest.SkipTest("needs the admin user and a company")
        cls.company = company
        client = appmod.app.test_client()
        with client.session_transaction() as s:
            s["_user_id"] = str(user)
            s["_fresh"] = True
            s["company_id"] = company
        import openpyxl
        cls.books = {}
        for vt in ("Receipt", "Payment", "Contra"):
            with contextlib.redirect_stdout(io.StringIO()):
                data = client.get("/download_voucher_template/" + vt).data
            cls.books[vt] = openpyxl.load_workbook(io.BytesIO(data))

    def rows(self, vt):
        ws = self.books[vt].worksheets[0]
        header = [c.value for c in ws[1]]
        out = []
        for row in ws.iter_rows(min_row=2, max_row=3, values_only=True):
            if any(v is not None for v in row):
                out.append(dict(zip(header, row)))
        return out

    def test_the_ledger_dropdown_follows_the_type_column(self):
        ws = self.books["Receipt"].worksheets[0]
        formulas = [dv.formula1 or "" for dv in ws.data_validations.dataValidation]
        self.assertTrue(any('INDIRECT(IF($F4="Credit","LEDGERS_CREDIT"' in f for f in formulas))

    def test_the_debit_list_is_the_allowed_list(self):
        book = self.books["Receipt"]
        ref = book.defined_names["LEDGERS_DEBIT"].attr_text
        sheet, cells = ref.split("!")
        start, end = cells.replace("$", "").split(":")
        ws = book[sheet.strip("'")]
        listed = [c.value for (c,) in ws[start:end]]
        self.assertEqual(sorted(listed),
                         M.allowed_ledger_names("Receipt", "Debit", company_id=self.company))

    def test_samples_are_balanced_pairs_dated_like_the_app(self):
        import re
        for vt in ("Receipt", "Payment", "Contra"):
            rows = self.rows(vt)
            self.assertEqual({r["Type"] for r in rows}, {"Debit", "Credit"}, vt)
            debit = sum(r["Amount"] for r in rows if r["Type"] == "Debit")
            credit = sum(r["Amount"] for r in rows if r["Type"] == "Credit")
            self.assertEqual(debit, credit, vt)
            for r in rows:
                self.assertRegex(str(r["Date"]), r"^\d{2}-\d{2}-\d{4}$", vt)

    def test_the_sample_cash_line_passes_the_validator(self):
        for vt in ("Receipt", "Contra"):
            debit = [r for r in self.rows(vt) if r["Type"] == "Debit"][0]
            ok, err = M.validate_voucher_ledger_groups(
                vt, [{"ledger_name": debit["Ledger Name"], "amount": debit["Amount"],
                      "type": "Debit"}], company_id=self.company)
            self.assertTrue(ok, (vt, err))


if __name__ == "__main__":
    unittest.main()
